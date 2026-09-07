# -*- coding: utf-8 -*-
"""掃描器的自我檢查。不需要任何真實展表、不需要 Google 金鑰，CI 的冒煙測試每次都跑。

檢查六件事：
  1. 數法對：拿手工造的分頁（四輪同名判定欄、Y/N/X/未測試/怪值、判 N 沒備註、★；
     以及第 5 列表頭＋一輪 AOS/iOS 兩欄的 binary 模板），用 openpyxl 走一遍，結果要跟手算的一樣。
  2. 兩條路同一個答案：同一張分頁改包成 Sheets API 的二維 values（GridSheet），
     再數一次，要跟 openpyxl 那條路一字不差。這是「兩種來源不會各算各的」的保證。
  3. 展表自報（測試進度分頁）的解析與比對。
  4. Google 雲端硬碟下載的 zip：不解壓直接找、直接讀，結果跟讀 xlsx 一樣。
  5. .gsheet 與雲端硬碟資料夾網址／ID 的解析。
  6. 任務（GitLab Issues）：展表標題對照鍵、類型推斷、Issue 解析。
任何一項不對就 exit 1。
"""
from __future__ import annotations
import io, json, os, sys, tempfile, zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="replace")
import openpyxl
import scan_progress as sp

fails = []


def check(name, got, want):
    if got != want:
        fails.append(f"✗ {name}\n    得到 {got!r}\n    應為 {want!r}")
    else:
        print(f"✓ {name}")


# ---------------------------------------------------------------- 1. 手工分頁
def make_book():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "手工分頁"
    ws.append(["第一次測試", None, None, None, "第二次測試"])          # r1
    ws.append(["所有項目數", "所有已測數", "所有未測數", "所有異常數"])   # r2
    ws.append(["=G3+H3+I3", "=COUNTIF(F5:F99,\"Y\")"])                 # r3（公式，沒快取值）
    ws.append(["項目", "ItemID", "描述", "企畫書對照", "備註",            # r4 表頭
               "確認內容正確", "異常狀況備註", "測試人員", "測試日期",      # 第一輪
               "確認內容正確", "測試備註", "測試人員", "測試日期",         # 第二輪（同名！）
               "確認內容正確", "測試備註", "測試人員", "測試日期",         # 第三輪
               "確認內容正確", "測試備註", "測試人員", "測試日期"])        # 第四輪
    rows = [
        # 名,   id,     描述, 對照,           備註, r1,     r1備註,    人, 日, r2,     ...
        ["甲", 1001,   "",  "",             "",  "Y",    "",       "a", "", "Y",   "", "", "", "未測試", "", "", "", "未測試"],
        ["乙", 1002,   "",  "★企畫書沒寫數量", "",  "N",    "數量不對", "a", "", "N",   "", "", "", "未測試"],
        ["丙", 1003,   "",  "",             "",  "N",    "",       "a", "", "未測試"],   # 判 N 沒備註
        ["丁", 1004,   "",  "",             "",  "X",    "",       "a", "", "X"],
        ["戊", 1005,   "",  "",             "",  "未測試", "",       "",  "", "未測試"],
        ["己", 1006,   "",  "",             "",  "ok",   "",       "a", "", "Y"],      # 怪值
        ["庚", 1007,   "",  "",             "",  "",     "",       "",  "", "Y"],      # 第一輪空白
        ["辛", 1008,   "",  "★ 要問特效",     "",  "Y",    "",       "a", "", "Y"],
    ]
    for r in rows:
        ws.append(r)
    ws2 = wb.create_sheet("測試進度")           # 應被跳過
    ws2.append(["整體", 99])
    ws3 = wb.create_sheet("圖片資料庫")          # 應被跳過
    ws3["F4"] = "確認內容正確"
    ws3["F5"] = "Y"
    ws4 = wb.create_sheet("沒有判定欄")          # 應回 None
    ws4["A4"] = "項目"
    ws4["A5"] = "x"
    return wb


WANT = {
    "分頁": "手工分頁",
    "計": {"Y": 2, "N": 2, "X": 1, "未測試": 1, "其他": 1},   # 只數第一輪 F 欄；庚是空白不算
    "總": 7, "母數": 5, "已測": 4,                              # 母數＝Y+N+未測試（X 與怪值不算）
    "表頭列": 4, "第一輪欄": ["F"],
    "star": [{"列": 6, "名": "乙", "id": "1002", "說明": "★企畫書沒寫數量"},
             {"列": 12, "名": "辛", "id": "1008", "說明": "★ 要問特效"}],
    "blank": [{"列": 7, "名": "丙", "id": "1003"}],
    "其他值": [{"列": 10, "值": "ok", "名": "己"}],
}

wb = make_book()
buf = io.BytesIO()
wb.save(buf)
buf.seek(0)
wb2 = openpyxl.load_workbook(buf, data_only=False)
got_xlsx = sp.scan_worksheets(wb2.worksheets)
check("openpyxl：只剩一張有判定欄的分頁", [s["分頁"] for s in got_xlsx], ["手工分頁"])
check("openpyxl：手算結果一致", got_xlsx[0], WANT)


# ------------------------------------------------- 2. 同一張表走 Sheets API 那條路
def to_api_values(ws):
    """模擬 Sheets API（UNFORMATTED_VALUE）回來的樣子：中間空格是 ""，列尾空格省略，
    整列空白是 []。公式格 API 會回算好的值；這裡沒有算，就給空字串（數法不看那幾格）。"""
    out = []
    for row in ws.iter_rows(values_only=True):
        vals = ["" if v is None else ("" if isinstance(v, str) and v.startswith("=") else v)
                for v in row]
        while vals and vals[-1] == "":
            vals.pop()
        out.append(vals)
    while out and out[-1] == []:
        out.pop()
    return out


grid_ws = [sp.GridSheet(ws.title, to_api_values(ws)) for ws in wb2.worksheets]
got_grid = sp.scan_worksheets(grid_ws)
check("GridSheet：跟 openpyxl 那條路一字不差", got_grid, got_xlsx)

# JSON 往返（Sheets API 回的是 JSON：整數不會帶 .0）
grid_ws_json = [sp.GridSheet(ws.title, json.loads(json.dumps(to_api_values(ws), ensure_ascii=False)))
                for ws in wb2.worksheets]
got_json = sp.scan_worksheets(grid_ws_json)
want_json = json.loads(json.dumps(got_xlsx, ensure_ascii=False))
got_ids = [x["id"] for x in got_json[0]["star"]]
check("GridSheet(JSON)：計數欄位一致", {k: got_json[0][k] for k in ("計", "總", "母數", "已測")},
      {k: want_json[0][k] for k in ("計", "總", "母數", "已測")})
check("GridSheet(JSON)：★ 與判N無備註列號一致",
      ([x["列"] for x in got_json[0]["star"]], [x["列"] for x in got_json[0]["blank"]]),
      ([x["列"] for x in want_json[0]["star"]], [x["列"] for x in want_json[0]["blank"]]))
check("GridSheet(JSON)：id 顯示（整數不帶 .0）", got_ids, ["1002", "1008"])


# ------------------------------ 2b. 第 5 列表頭 ＋ 一輪兩欄（AOS／iOS）＋ 各欄自己的備註欄
def make_binary_book():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "binary基本測試"
    #            A          B  C  D  E   F            G          H  I   J            K          L  M   N(第二輪)
    ws.append(["登入口帳號", "", "", "", "", "第一次測試", "", "", "", "", "", "", "", "第二次測試"])   # r1 標籤
    ws.append(["所有項目數", "所有已測數", "所有未測數", "所有異常數"])                                  # r2
    ws.append(["=G3+H3+I3", "=COUNTIF(F6:M99,\"Y\")", "=COUNTIF(F6:M99,\"未測試\")", "=COUNTIF(F6:M99,\"N\")"])  # r3
    ws.append(["binary基本測試", "", "", "", "", "ＡＯＳ", "", "", "", "IＯＳ", "", "", "", "ＡＯＳ", "", "", "", "IＯＳ"])  # r4 群組標題
    ws.append(["項目", "說明", "描述", "測試項目", "備註",
               "確認內容正確", "測試備註", "測試人員", "測試日期",     # F..I  AOS 第一輪
               "確認內容正確", "測試備註", "測試人員", "測試日期",     # J..M  iOS 第一輪
               "確認內容正確", "測試備註", "測試人員", "測試日期",     # N..Q  AOS 第二輪
               "確認內容正確", "測試備註", "測試人員", "測試日期"])    # R..U  iOS 第二輪
    ws.append(["登入", "", "", "", "", "Y", "", "a", "", "Y", "", "a", "", "Y", "", "", "", "Y"])
    ws.append(["建角", "", "", "", "", "N", "閃退", "a", "", "Y", "", "a", "", "未測試"])          # AOS N 有備註
    ws.append(["刪角", "", "", "", "", "Y", "", "a", "", "N", "", "a", "", "未測試"])              # iOS N 沒備註（K 空）
    ws.append(["儲值", "", "", "", "", "未測試", "", "", "", "X", "", "", "", "未測試"])
    return wb

WANT_BIN = {
    "分頁": "binary基本測試",
    "計": {"Y": 4, "N": 2, "X": 1, "未測試": 1, "其他": 0},   # F 與 J 兩欄都算；N、R 是第二輪不算
    "總": 8, "母數": 7, "已測": 6,
    "表頭列": 5, "第一輪欄": ["F", "J"],
    "star": [], "blank": [{"列": 8, "名": "刪角", "id": ""}], "其他值": [],
}
wbb = make_binary_book()
bufb = io.BytesIO(); wbb.save(bufb); bufb.seek(0)
wbb2 = openpyxl.load_workbook(bufb, data_only=False)
got_bin = sp.scan_worksheets(wbb2.worksheets)
check("第 5 列表頭＋雙欄輪次：手算結果一致", got_bin[0] if got_bin else None, WANT_BIN)
got_bin_grid = sp.scan_worksheets([sp.GridSheet(ws.title, to_api_values(ws)) for ws in wbb2.worksheets])
check("第 5 列表頭＋雙欄輪次：GridSheet 同答案", got_bin_grid, got_bin)

# 沒有「第N次測試」標籤、但分頁自己的 COUNTIF(F6:M99) 還在：照 COUNTIF 的欄範圍，F 與 J 都算
wbc = make_binary_book(); wsc = wbc.active
for c in wsc[1]: c.value = None
bufc = io.BytesIO(); wbc.save(bufc); bufc.seek(0)
got_c = sp.scan_worksheets(openpyxl.load_workbook(bufc).worksheets)
check("沒標籤但有 COUNTIF 欄範圍：第一輪 F、J", (got_c[0]["第一輪欄"], got_c[0]["計"]["Y"]), (["F", "J"], 4))
# 標籤跟 COUNTIF 都沒有：退回連續法（只剩 F，J 被中間的備註欄隔開）
for c in wsc[3]: c.value = None
bufc = io.BytesIO(); wbc.save(bufc); bufc.seek(0)
got_c = sp.scan_worksheets(openpyxl.load_workbook(bufc).worksheets)
check("沒標籤也沒 COUNTIF：退回連續法只剩 F", (got_c[0]["第一輪欄"], got_c[0]["計"]["Y"]), (["F"], 2))
check("COUNTIF 欄範圍解析", sp.round1_bounds([{2: "所有已測數"}, {2: '=COUNTIF(Q6:X167,"Y")'}]), (17, 24))
check("COUNTIF 欄範圍解析：沒有公式", sp.round1_bounds([{2: "所有已測數"}, {2: 285}]), None)

# 2b-2. 製作清單型：五個判定欄，其中「機率正確」「payback正確」沒寫「確認」；表頭在第 5 列、標籤只有一個「第1次」
def make_craft_book():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "製作清單(共同)(一般)"
    ws.append([None] * 16 + ["所有項目數", "所有已測數", "所有未測數", "所有異常數"])               # r1
    ws.append([None] * 16 + ["=R2+S2+T2", '=COUNTIF(Q6:X167,"Y")', '=COUNTIF(Q6:X167,"未測試")', '=COUNTIF(Q6:X167,"N")'])  # r2
    ws.append([None] * 16 + ["第1次"])                                                             # r3
    ws.append(["[共同] 月初/限定 BM"])                                                              # r4 說明
    ws.append(["CraftID", "詳細區分", "伺服器", "分類", "等級", "ItemID", "製作結果道具", "材料", "數量", "限制單位",
               "限制數量", "製作機率", "實際返還道具", "返還道具個數", "血盟返還", "週次別",
               "確認製作時間區間設定正確", "確認伺服器限制與製作限制正確", "確認材料與結果物與數量設定正確", "機率正確", "payback正確",
               "測試備註", "測試人員", "測試日期"])                                                    # r5 表頭
    ws.append([1, "", "", "", "", 11, "甲"] + [""] * 9 + ["Y", "Y", "Y", "Y", "Y", "", "a", ""])
    ws.append([2, "", "", "", "", 12, "乙"] + [""] * 9 + ["Y", "N", "Y", "X", "X", "", "a", ""])       # N 沒備註
    return wb
wbk = make_craft_book(); bufk = io.BytesIO(); wbk.save(bufk); bufk.seek(0)
got_k = sp.scan_worksheets(openpyxl.load_workbook(bufk).worksheets)[0]
check("製作清單型：五欄都算、一格一項", (got_k["表頭列"], got_k["第一輪欄"], got_k["計"], got_k["母數"]),
      (5, ["Q", "R", "S", "T", "U"], {"Y": 7, "N": 1, "X": 2, "未測試": 0, "其他": 0}, 8))
check("製作清單型：判 N 沒備註用最右邊那欄備註", got_k["blank"], [{"列": 7, "名": "乙", "id": "12"}])

# 2b-3. 變身確認型：說明文字寫在表頭上方、左右兩張表、中段重複表頭
def make_morph_book():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "變身確認"
    ws.append([None] * 5 + ["所有項目數", "所有已測數", "所有未測數", "所有異常數"])                       # r1
    ws.append(["更新時間：x"] + [None] * 4 + ["=G2+H2+I2", '=COUNTIF(F3:P185,"Y")', '=COUNTIF(F3:P185,"未測試")', '=COUNTIF(F3:P185,"N")'])
    ws.append(["死神變身", "確認常發生問題的變身，有時間再抽查其他變身", "", "", "", "/1 887 100", "", "", "", "", "槍手", "確認常發生問題的變身，有時間再抽查其他變身"])
    ws.append(["階級-變身", "測試項目", "", "", "備註", "確認內容正確", "測試備註", "測試人員", "測試日期", "", "階級-變身", "測試項目", "", "", "備註", "確認內容正確", "測試備註", "測試人員", "測試日期"])
    ws.append(["英雄-甲", "變身顯示正常", "", "", 1, "Y", "", "a", "", "", "英雄-乙", "變身顯示正常", "", "", 2, "Y", "", "a", ""])
    ws.append([])
    ws.append(["雷神變身", "確認常發生問題的變身，有時間再抽查其他變身", "", "", "", "/1 157591 1", "", "", "", "", "龍鬥士", "確認常發生問題的變身，有時間再抽查其他變身"])
    ws.append(["階級-變身", "測試項目", "", "", "備註", "確認內容正確", "測試備註", "測試人員", "測試日期", "", "階級-變身", "測試項目", "", "", "備註", "確認內容正確", "測試備註", "測試人員", "測試日期"])
    ws.append(["英雄-丙", "變身顯示正常", "", "", 3, "N", "掉幀", "a", "", "", "英雄-丁", "變身顯示正常", "", "", 4, "Y", "", "a", ""])
    return wb
wbm = make_morph_book(); bufm = io.BytesIO(); wbm.save(bufm); bufm.seek(0)
got_m = sp.scan_worksheets(openpyxl.load_workbook(bufm).worksheets)[0]
check("變身確認型：表頭是第 4 列（不是第 3 列的說明文字）、左右兩欄都算、重複表頭與 /指令 不算判定值",
      (got_m["表頭列"], got_m["第一輪欄"], got_m["計"]), (4, ["F", "P"], {"Y": 3, "N": 1, "X": 0, "未測試": 0, "其他": 0}))

# 2b-4. 掃描範圍照「測試進度」公式
check("掃描範圍解析", sp.scope_from_formulas([
    ["第一次測試", None, None, None, "第二次測試"],
    ["所有項目數", "所有已測數", "所有未測數", "所有異常數"],
    ["='道具說明使用'!AB3+'序號兌換測試'!G3", "='道具說明使用'!AC3+'序號兌換測試'!H3", "='怪物設定(機率&掉落物)'!Q3+binary基本測試!H3", "=0"]]),
    ["道具說明使用", "序號兌換測試", "怪物設定(機率&掉落物)", "binary基本測試"])
check("掃描範圍解析：沒有公式（只有值）就空", sp.scope_from_formulas([["第一次測試"], ["所有項目數", "所有已測數"], [95, 79]]), [])
wbs = openpyxl.Workbook(); s1 = wbs.active; s1.title = "有算"; s2 = wbs.create_sheet("沒掛"); s3 = wbs.create_sheet("測試進度")
for s in (s1, s2):
    s.append(["第一次測試"]); s.append(["所有項目數", "所有已測數"]); s.append(["=B3", '=COUNTIF(B5:B99,"Y")'])
    s.append(["項目", "確認內容正確", "異常狀況備註", "測試人員"]); s.append(["甲", "Y", "", "a"]); s.append(["乙", "N", "x", "a"])
s3.append(["第一次測試"]); s3.append(["所有項目數", "所有已測數", "所有未測數", "所有異常數"]); s3.append(["='有算'!A3+'不存在的分頁'!A3", "='有算'!B3", "=0", "=0"])
bufs = io.BytesIO(); wbs.save(bufs); bufs.seek(0)
wbs2 = openpyxl.load_workbook(bufs)
formulas = {ws.title: [list(r) for r in ws.iter_rows(min_row=1, max_row=15, values_only=True)] for ws in wbs2.worksheets}
sheets_s, info_s = sp.scan_book(wbs2.worksheets, formulas)
check("掃描範圍：只算公式點到的分頁", [s["分頁"] for s in sheets_s], ["有算"])
check("掃描範圍：沒掛進公式的分頁列出來、公式點到但不存在的分頁列出來",
      (info_s["範圍"], info_s["未掛進進度表"], info_s["進度表點到但沒這分頁"], info_s["進度表點到但掃不到"]),
      ("進度表", [{"分頁": "沒掛", "母數": 2, "已測": 2}], ["不存在的分頁"], []))
sheets_a, info_a = sp.scan_book(wbs2.worksheets, {"測試進度": [["第一次測試"], ["所有項目數"], [4]]})
check("掃描範圍：公式抓不到就全掃", ([s["分頁"] for s in sheets_a], info_a["範圍"]), (["有算", "沒掛"], "自動"))
# 公式點到的分頁一張都對不到（分頁改名了）：不能讓整本消失，退回全掃並註記
formulas_renamed = dict(formulas); formulas_renamed["測試進度"] = [["第一次測試"], ["所有項目數", "所有已測數"], ["='舊名字'!A3", "='舊名字'!B3"]]
sheets_r, info_r = sp.scan_book(wbs2.worksheets, formulas_renamed)
check("掃描範圍：公式全對不到就退回全掃、不吞掉整本",
      ([s["分頁"] for s in sheets_r], info_r["範圍"].startswith("自動（"), info_r["進度表點到但沒這分頁"]),
      (["有算", "沒掛"], True, ["舊名字"]))
# 同一本走 Sheets API 的形狀（值：""、尾巴截斷；公式：FORMULA render 的二維 list）要跟 openpyxl 同答案
grids = [sp.GridSheet(ws.title, to_api_values(ws)) for ws in wbs2.worksheets]
def api_formula_shape(rows):
    out = []
    for r in rows:
        vals = ["" if v is None else v for v in r]
        while vals and vals[-1] == "": vals.pop()
        out.append(vals)
    while out and out[-1] == []: out.pop()
    return out
formulas_api = {t: api_formula_shape(r) for t, r in formulas.items()}
check("掃描範圍：GridSheet＋API 形狀的公式，跟 openpyxl 同答案", sp.scan_book(grids, formulas_api), (sheets_s, info_s))
check("COUNTIF 欄範圍解析：多段相加取聯集", sp.round1_bounds([{2: "所有已測數"}, {2: '=COUNTIF(F6:F99,"Y")+COUNTIF(J6:J99,"Y")'}]), (6, 10))
check("COUNTIF 欄範圍解析：點到別張分頁就不信", sp.round1_bounds([{2: "所有已測數"}, {2: "=COUNTIF('別張'!F6:M99,\"Y\")"}]), None)
check("掃描範圍解析：字串串接不會誤抓", sp.scope_from_formulas([[], ["所有項目數"], ['="x"&Sheet1!A1+SUM(A1:A3)']]), ["Sheet1"])

# ------------------------------------------------ 2c. 展表自報（測試進度）解析與比對
rows = [["第一次測試", None, None, None, "第二次測試"],
        ["所有項目數", "所有已測數", "所有未測數", "所有異常數", "所有項目數"],
        [17, 16.0, 0, 1, 17]]
check("自報解析", sp.self_report(rows), {"項目": 17, "Y": 16, "未測試": 0, "N": 1})
check("自報解析：API 省略尾巴", sp.self_report([[], ["所有項目數", "所有已測數", "所有未測數", "所有異常數"], [7, 7]]), None)
check("自報解析：沒有標籤", sp.self_report([["整體", 3]]), None)
agg = {"Y": 16, "N": 1, "X": 2, "未測試": 0, "其他": 3}
check("自報比對：一致就沒有差異（X 與怪值不算進項目數）", sp.compare_report({"項目": 17, "Y": 16, "未測試": 0, "N": 1}, agg), [])
check("自報比對：差在哪一欄", sp.compare_report({"項目": 20, "Y": 16, "未測試": 3, "N": 1}, agg),
      [{"欄位": "項目數", "自報": 20, "看板": 17}, {"欄位": "未測試", "自報": 3, "看板": 0}])


# ------------------------------------------------ 2d. Google 雲端硬碟下載的 zip（不解壓直接讀）
with tempfile.TemporaryDirectory() as td:
    zp = os.path.join(td, "drive-download-test.zip")
    with zipfile.ZipFile(zp, "w") as z:
        z.writestr("0826(binary)(測)/QA - 手工.xlsx", buf.getvalue())            # 展表（在子資料夾）
        z.writestr("0826(binary)(測)/@BM/QA - binary.xlsx", bufb.getvalue())    # 展表（更深一層）
        z.writestr("0826(binary)(測)/企畫書 的副本.xlsx", buf.getvalue())         # 非前綴，要略過
        z.writestr("__MACOSX/._QA - 手工.xlsx", b"junk")                       # Mac 垃圾，要略過
    st = {"候選": 0, "非前綴略過": 0, "讀不開": 0, "沒有判定欄": 0, "雲端xlsx未讀": 0, "無自報可對": 0}
    found = sp.discover_zip(zp, "QA - ", st)
    check("zip：找到的展表", [n for n, _, _ in found], ["QA - binary.xlsx", "QA - 手工.xlsx"])
    check("zip：候選／略過統計", (st["候選"], st["非前綴略過"]), (3, 1))
    check("zip：最上層資料夾名（給版本檢查）", sp.zip_top_folder(zp), "0826(binary)(測)")
    got_zip = {n: sp.read_book(k, w, None)[0] for n, k, w in found}
    check("zip：讀出來跟直接讀 xlsx 一字不差", (got_zip["QA - 手工.xlsx"], got_zip["QA - binary.xlsx"]), (got_xlsx, got_bin))
    zp2 = os.path.join(td, "flat.zip")
    with zipfile.ZipFile(zp2, "w") as z:
        z.writestr("QA - a.xlsx", buf.getvalue()); z.writestr("QA - b.xlsx", bufb.getvalue())
    check("zip：檔案直接在根目錄就沒有資料夾名", sp.zip_top_folder(zp2), None)


# ------------------------------------------------ 2e. 任務（GitLab Issues）對照鍵與解析
import fetch_issues as fi
check("task_key：去 QA- 前綴、全半形、斜線底線空白都抹掉",
      [fi.task_key("QA - 血盟/伺服器移民_20260826"), fi.task_key("血盟 伺服器移民 20260826"), fi.task_key("ＲＣ２ binary更新基本測試"), fi.task_key("QA-rc2binary更新基本測試")],
      ["血盟伺服器移民20260826", "血盟伺服器移民20260826", "rc2binary更新基本測試", "rc2binary更新基本測試"])
check("guess_type", [fi.guess_type(t) for t in ("活動_共同_中元節", "(台灣提案)活動_哈爾巴斯", "BM_月底_x", "IAP_共同_x", "系統_重生伺服器", "RC2 binary更新基本測試", "QA - 送審基本測試")],
      ["活動", "活動", "BM", "BM", "系統", "基本測試", "基本測試"])
parsed = fi.parse_issues([
    {"iid": 3, "title": "QA - 活動_共同_神秘的流浪馬戲團", "state": "opened", "labels": ["狀態::測試中", "輪次::第2輪", "類型::活動"],
     "assignees": [{"name": "王小明", "username": "xm", "avatar_url": "https://x/a.png"}], "milestone": {"title": "0826"},
     "web_url": "https://g/x/-/issues/3", "updated_at": "2026-09-04T01:02:03.000Z", "due_date": None},
    {"iid": 4, "title": "送審基本測試", "state": "closed", "labels": [], "assignees": [], "milestone": None,
     "web_url": "https://g/x/-/issues/4", "updated_at": "2026-09-04T01:02:03.000Z"},
])
check("parse_issues：標籤拆成狀態／輪次／類型、負責人、里程碑",
      {k: parsed[0][k] for k in ("鍵", "狀態", "輪次", "類型", "里程碑", "open")} | {"負責人": parsed[0]["負責人"][0]["名"]},
      {"鍵": "活動共同神秘的流浪馬戲團", "狀態": "測試中", "輪次": "第2輪", "類型": "活動", "里程碑": "0826", "open": True, "負責人": "王小明"})
check("parse_issues：沒標籤的關閉卡＝完成、類型用標題猜", (parsed[1]["狀態"], parsed[1]["輪次"], parsed[1]["類型"], parsed[1]["open"]),
      ("完成", None, "基本測試", False))

# ---------------------------------------------------------------- 3. 解析
with tempfile.TemporaryDirectory() as td:
    p = os.path.join(td, "a.gsheet")
    open(p, "w", encoding="utf-8").write(json.dumps(
        {"url": "https://docs.google.com/spreadsheets/d/1AbC_dEf-123/edit?usp=drivesdk",
         "doc_id": "1AbC_dEf-123", "email": "x@y.z", "resource_id": "spreadsheet:1AbC_dEf-123"}))
    check(".gsheet 解析（doc_id）", sp.gsheet_id(p), "1AbC_dEf-123")
    open(p, "w", encoding="utf-8").write(json.dumps(
        {"url": "https://docs.google.com/spreadsheets/d/1XyZ/edit"}))
    check(".gsheet 解析（只有 url）", sp.gsheet_id(p), "1XyZ")
    open(p, "w", encoding="utf-8").write("garbage")
    check(".gsheet 解析（壞檔回 None）", sp.gsheet_id(p), None)

check("資料夾網址解析", sp.drive_folder_id("https://drive.google.com/drive/u/0/folders/1Folder_ID-xx?usp=sharing"),
      "1Folder_ID-xx")
check("資料夾 ID 解析", sp.drive_folder_id("1AbCdEfGhIjKlMnOpQrStUvWxYz"), "1AbCdEfGhIjKlMnOpQrStUvWxYz")
check("本機路徑不是資料夾 ID", sp.drive_folder_id(r"G:\共用雲端硬碟\E_天堂M(版控專用)"), None)
check("檔名前綴過濾", [sp.wanted(n, "QA - ") for n in ("QA - 甲.xlsx", "~$QA - 甲.xlsx", "企畫書.xlsx", "提問清單.xlsx")],
      [True, False, False, False])

if fails:
    print("\n".join(fails), file=sys.stderr)
    sys.exit(f"selftest 失敗 {len(fails)} 項")
print("selftest OK")
